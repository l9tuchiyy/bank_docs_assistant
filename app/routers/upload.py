import io
from pathlib import Path
from uuid import UUID, uuid4

from fastapi import APIRouter, File, Form, HTTPException, Query, UploadFile, status

from app.config import ALLOWED_EXTENSIONS, MAX_FILE_SIZE, UPLOAD_DIR
from app.deps import engine
from app.services import audit
from app.schemas import DeleteResponseSchema, UploadResponseSchema

router = APIRouter(prefix="/api", tags=["Документы"])


def extract_text(ext: str, content: bytes) -> str:
    if ext == ".docx":
        import docx
        d = docx.Document(io.BytesIO(content))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            parts += [" | ".join(c.text for c in r.cells) for r in t.rows]
        return "\n".join(parts)
    if ext == ".pdf":
        import pypdf
        return "\n".join((p.extract_text() or "") for p in pypdf.PdfReader(io.BytesIO(content)).pages)
    return content.decode("utf-8", errors="ignore")


@router.post("/upload", response_model=UploadResponseSchema, status_code=status.HTTP_201_CREATED, summary="Загрузить проект документа для проверки")
async def upload_document(file: UploadFile = File(...), session_id: str | None = Form(default=None)) -> UploadResponseSchema:
    filename = file.filename or ""
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(400, "Можно загрузить только PDF, DOCX, TXT или MD")
    content = await file.read()
    if not content:
        raise HTTPException(400, "Файл пустой")
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(400, "Максимальный размер — 10 МБ")
    try:
        text = extract_text(ext, content)
    except Exception:
        raise HTTPException(400, "Не удалось прочитать файл")
    if not text.strip():
        raise HTTPException(400, "В файле не найден текст (скан без распознавания?)")
    doc_id = str(uuid4())
    UPLOAD_DIR.mkdir(exist_ok=True)
    (UPLOAD_DIR / f"{doc_id}{ext}").write_bytes(content)
    sid = engine.attach_document(session_id, filename, text, doc_id)
    return UploadResponseSchema(document_id=doc_id, filename=filename, status="uploaded", session_id=sid, chars=len(text))


@router.delete("/upload/{document_id}", response_model=DeleteResponseSchema, summary="Удалить загруженный документ")
def delete_document(document_id: str, session_id: str = Query(min_length=1)) -> DeleteResponseSchema:
    """Удаляет документ из сессии и файл с диска. Удалить может только сессия, которая его загрузила."""
    try:
        doc_id = str(UUID(document_id))          # защита от path traversal: принимаем только UUID
    except ValueError:
        raise HTTPException(400, "Некорректный идентификатор документа")
    if not engine.detach_document(session_id, doc_id):
        raise HTTPException(404, "Документ не найден в этой сессии")
    for f in UPLOAD_DIR.glob(f"{doc_id}.*"):
        f.unlink(missing_ok=True)
    audit.log(session_id, "employee", "doc_deleted", [])
    return DeleteResponseSchema(document_id=doc_id, status="deleted", session_id=session_id)
